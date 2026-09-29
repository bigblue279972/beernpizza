"""
Build the betting ledger workbook to Cliff's own column spec.

His ten columns, in his order and under his names:
  Opening Balance, Competition, Teams, Bet Type, Odds Required, Betfair Odds,
  Bet Size, Win/Lose, Total Won/Lost, Closing Balance

A Date column goes in front (the balance graph needs one) and a few extras sit
after the ten, so the spec he gave is untouched and everything added is clearly
additional. Back/Lay is one of those extras and is treated as Back when blank,
so a normal bet needs nothing typed in it.

The pairing of Odds Required against Betfair Odds is the useful thing here: the
first is what he decided he needed, the second is what he got, so the gap is a
recorded claim about edge that the Dashboard can later test against what
actually happened.

Usage:  python build_bet_ledger.py <output.xlsx> [--demo]
"""
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side as BSide
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment

LAST = int(os.environ.get("LEDGER_ROWS", "1001"))
L1, LN = 14, 60                     # Settings list rows
SPARE = 12

FONT = "Arial"
NAVY, TEAL, GREY = "1F3A5F", "0F6E6E", "F2F4F7"
AMBER, GREEN, RED, YELLOW = "FFF3CD", "E3F4E6", "FBE3E3", "FFF2A8"
MONEY = '$#,##0.00;[Red]($#,##0.00)'
MONEY0 = '$#,##0;[Red]($#,##0)'
PCT = '0.0%'
PCTS = '+0.0%;-0.0%;0.0%'
ODDS = '0.00'
DATEF = 'DD/MM/YYYY'
INT = '#,##0'

thin = BSide(style="thin", color="D0D5DD")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def title_cell(ws, ref, text, size=14, colour=NAVY):
    ws[ref] = text
    ws[ref].font = Font(name=FONT, size=size, bold=True, color=colour)


def section(ws, ref, text):
    ws[ref] = text
    ws[ref].font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
    ws[ref].fill = PatternFill("solid", fgColor=TEAL)
    ws[ref].alignment = Alignment(horizontal="left", vertical="center")


def label(ws, ref, text, bold=False, italic=False, colour="344054"):
    ws[ref] = text
    ws[ref].font = Font(name=FONT, size=10, bold=bold, italic=italic, color=colour)


def value(ws, ref, formula, fmt=None):
    ws[ref] = formula
    ws[ref].font = Font(name=FONT, size=11, bold=True, color=NAVY)
    ws[ref].alignment = Alignment(horizontal="right")
    if fmt:
        ws[ref].number_format = fmt


def yellow(ws, ref, fmt, note):
    c = ws[ref]
    c.fill = PatternFill("solid", fgColor=YELLOW)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    c.font = Font(name=FONT, size=11, bold=True)
    c.alignment = Alignment(horizontal="center")
    c.comment = Comment(note, "Betting Ledger")


# key, header, width, format, kind ("in" typed / "auto" calculated / "help" hidden)
COLUMNS = [
    ("date",     "Date",             11, DATEF,  "in"),
    ("open_bal", "Opening Balance",  15, MONEY,  "auto"),
    ("comp",     "Competition",      18, None,   "in"),
    ("teams",    "Teams",            26, None,   "in"),
    ("bet_type", "Bet Type",         20, None,   "in"),
    ("req",      "Odds Required",    13, ODDS,   "in"),
    ("got",      "Betfair Odds",     12, ODDS,   "in"),
    ("size",     "Bet Size",         11, MONEY0, "in"),
    ("result",   "Win/Lose",         11, None,   "in"),
    ("pl",       "Total Won/Lost",   14, MONEY,  "auto"),
    ("close_bal", "Closing Balance", 15, MONEY,  "auto"),
    ("side",     "Back/Lay",         10, None,   "in"),
    ("edge",     "Edge %",           10, PCTS,   "auto"),
    ("close_od", "Closing Odds",     12, ODDS,   "in"),
    ("clv",      "CLV %",            10, PCTS,   "auto"),
    ("notes",    "Notes",            28, None,   "in"),
    ("peak",     "Peak",             12, MONEY,  "help"),
    ("dd",       "Drawdown",         12, MONEY,  "help"),
    ("band",     "Odds Band",        12, None,   "help"),
    ("month",    "Month",            10, None,   "help"),
]
COL = {k: get_column_letter(i) for i, (k, *_r) in enumerate(COLUMNS, start=1)}
IDX = {k: i for i, (k, *_r) in enumerate(COLUMNS, start=1)}


def c(key, row):
    return f"${COL[key]}{row}"


def rng(key):
    return f"${COL[key]}$2:${COL[key]}${LAST}"


def lg(key):
    return f"'Bet Log'!{rng(key)}"


wb = Workbook()
ws = wb.active
ws.title = "Bet Log"
wd = wb.create_sheet("Dashboard")
wk = wb.create_sheet("Breakdowns")
wset = wb.create_sheet("Settings")
wh = wb.create_sheet("How To Use")
S = "Settings!"

# ================================================================ SETTINGS
title_cell(wset, "A1", "SETTINGS")
label(wset, "A2", "Fill the two yellow cells. Everything else on this sheet is a "
                  "dropdown list you can edit.", italic=True)
for i, (lab, ref, fmt, note) in enumerate([
        ("Starting balance ($)", "B3", MONEY0,
         "What your Betfair balance was when you started this log. "
         "Every Opening and Closing Balance builds from this one number."),
        ("Betfair commission", "B4", PCT,
         "Your actual rate. Type 5% or 0.05. It comes off winnings only, "
         "never off a losing bet and never off the odds.")]):
    label(wset, f"A{3+i}", lab, bold=True)
    yellow(wset, ref, fmt, note)
    label(wset, f"C{3+i}", note, italic=True, colour="667085")

section(wset, "A11", "DROPDOWN LISTS - add a row under any list and the dropdown picks it up.")
wset.merge_cells("A11:J11")

lists = {
    "D": ("Competition", [
        "EPL", "Championship", "League One", "League Two", "Scottish Premiership",
        "La Liga", "Serie A", "Bundesliga", "Ligue 1", "Eredivisie", "Primeira Liga",
        "Belgian Pro League", "A-League Men", "MLS", "Liga MX", "Brasileirao",
        "Argentine Primera", "J1 League", "Chinese Super League", "K League 1",
        "UEFA Champions League", "Europa League", "International", "AFL", "Other"]),
    "E": ("Bet Type", [
        "Match Odds", "Double Chance", "Draw No Bet", "Both Teams To Score",
        "Over/Under 0.5 Goals", "Over/Under 1.5 Goals", "Over/Under 2.5 Goals",
        "Over/Under 3.5 Goals", "Over/Under 4.5 Goals", "Over/Under 5.5 Goals",
        "Asian Handicap", "Handicap", "Line", "Match Total", "Total Points",
        "Correct Score", "Half Time / Full Time", "First Goalscorer", "Other"]),
    "F": ("Win/Lose", ["Win", "Lose", "Half Win", "Half Lose", "Void"]),
    "G": ("Back/Lay", ["Back", "Lay"]),
    "H": ("Odds Band", ["1.01 - 1.49", "1.50 - 1.99", "2.00 - 2.99",
                        "3.00 - 4.99", "5.00 - 9.99", "10.00 +"]),
}
RANGE = {}
for col, (hdr, items) in lists.items():
    h = wset[f"{col}{L1}"]
    h.value = hdr
    h.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor=NAVY)
    h.alignment = Alignment(horizontal="center")
    for i, v in enumerate(items):
        cell = wset[f"{col}{L1+1+i}"]
        cell.value = v
        cell.font = Font(name=FONT, size=10)
    RANGE[col] = f"Settings!${col}${L1+1}:${col}${L1+len(items)+SPARE}"
    for r in range(L1 + 1 + len(items), L1 + 1 + len(items) + SPARE):
        cell = wset[f"{col}{r}"]
        cell.fill = PatternFill("solid", fgColor="F7F9FC")
        cell.border = BOX
for col, w in {"A": 26, "B": 14, "C": 60, "D": 24, "E": 24,
               "F": 12, "G": 11, "H": 14}.items():
    wset.column_dimensions[col].width = w

# ================================================================ BET LOG
HDR = {"in": NAVY, "auto": "475467", "help": "98A2B3"}
for i, (key, name, width, fmt, kind) in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=i, value=name)
    cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=HDR[kind])
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BOX
    ws.column_dimensions[get_column_letter(i)].width = width
    if kind == "help":
        ws.column_dimensions[get_column_letter(i)].hidden = True
ws.row_dimensions[1].height = 30
ws.freeze_panes = "B2"

for key, text in {
    "date": "Press Ctrl+; to type today's date.\nDark blue headers = you type these.\n"
            "Grey headers = calculated, never type in them.",
    "open_bal": "Your balance before this bet. Comes from the closing balance of the "
                "row above, so the column chains all the way back to the starting "
                "balance on Settings.",
    "req": "The price you decided you needed before you looked at the market.",
    "got": "The price you actually took on Betfair. The gap between this and Odds "
           "Required is your claimed edge, and the Dashboard tests it against what "
           "really happened.",
    "result": "Did YOUR BET win - not did the team win. A lay that lands is a Win.\n"
              "Half Win and Half Lose are for Asian lines.",
    "pl": "Profit or loss on this bet, with Betfair commission already taken off "
          "any winnings.",
    "side": "Leave blank for a normal back bet. Only fill it in when you laid.",
    "close_od": "Optional. The price the market settled at, for CLV. "
                "Leave blank if you are not capturing closing prices yet.",
}.items():
    ws[f"{COL[key]}1"].comment = Comment(text, "Betting Ledger")

for r in range(2, LAST + 1):
    p = r - 1
    lay = f'{c("side", r)}="Lay"'
    ws[f"{COL['open_bal']}{r}"] = (f'={S}$B$3' if r == 2 else f'={c("close_bal", p)}')
    ws[f"{COL['pl']}{r}"] = (
        f'=IF({c("result", r)}="","",'
        f'IF({c("result", r)}="Void",0,'
        f'IF(OR({c("size", r)}="",{c("got", r)}=""),"",'
        f'IF({lay},'
        f'IF({c("result", r)}="Win",{c("size", r)}*(1-{S}$B$4),'
        f'IF({c("result", r)}="Half Win",{c("size", r)}*0.5*(1-{S}$B$4),'
        f'IF({c("result", r)}="Lose",-{c("size", r)}*({c("got", r)}-1),'
        f'-{c("size", r)}*({c("got", r)}-1)*0.5))),'
        f'IF({c("result", r)}="Win",{c("size", r)}*({c("got", r)}-1)*(1-{S}$B$4),'
        f'IF({c("result", r)}="Half Win",'
        f'{c("size", r)}*({c("got", r)}-1)*0.5*(1-{S}$B$4),'
        f'IF({c("result", r)}="Lose",-{c("size", r)},-{c("size", r)}*0.5)))))))')
    ws[f"{COL['close_bal']}{r}"] = (
        f'={c("open_bal", r)}+IF(ISNUMBER({c("pl", r)}),{c("pl", r)},0)')
    ws[f"{COL['edge']}{r}"] = (
        f'=IF(OR({c("req", r)}="",{c("got", r)}=""),"",'
        f'IFERROR(IF({lay},1-{c("got", r)}/{c("req", r)},'
        f'{c("got", r)}/{c("req", r)}-1),""))')
    ws[f"{COL['clv']}{r}"] = (
        f'=IF(OR({c("close_od", r)}="",{c("got", r)}=""),"",'
        f'IFERROR(IF({lay},1-{c("got", r)}/{c("close_od", r)},'
        f'{c("got", r)}/{c("close_od", r)}-1),""))')
    ws[f"{COL['peak']}{r}"] = (f'={c("close_bal", r)}' if r == 2
                               else f'=MAX({c("peak", p)},{c("close_bal", r)})')
    ws[f"{COL['dd']}{r}"] = f'={c("peak", r)}-{c("close_bal", r)}'
    ws[f"{COL['band']}{r}"] = (
        f'=IF({c("got", r)}="","",IF({c("got", r)}<1.5,"1.01 - 1.49",'
        f'IF({c("got", r)}<2,"1.50 - 1.99",IF({c("got", r)}<3,"2.00 - 2.99",'
        f'IF({c("got", r)}<5,"3.00 - 4.99",IF({c("got", r)}<10,"5.00 - 9.99",'
        f'"10.00 +"))))))')
    ws[f"{COL['month']}{r}"] = (
        f'=IF({c("date", r)}="","",TEXT({c("date", r)},"YYYY-MM"))')
    for i, (key, _n, _w, fmt, kind) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=r, column=i)
        cell.font = Font(name=FONT, size=10)
        if fmt:
            cell.number_format = fmt
        cell.fill = PatternFill("solid", fgColor={"in": "FFFFFF", "auto": GREY,
                                                  "help": GREY}[kind])
    for key in ("result", "side", "band"):
        ws[f"{COL[key]}{r}"].alignment = Alignment(horizontal="center")

for key, src in (("comp", "D"), ("bet_type", "E"), ("result", "F"), ("side", "G")):
    dv = DataValidation(type="list", formula1=RANGE[src], allow_blank=True)
    dv.error = "Pick from the list, or add your entry on the Settings tab first."
    dv.errorTitle = "Not on the list"
    ws.add_data_validation(dv)
    dv.add(f"{COL[key]}2:{COL[key]}{LAST}")

CF = ws.conditional_formatting
CF.add(f"{COL['result']}2:{COL['result']}{LAST}", FormulaRule(
    formula=[f'AND({c("date", 2)}<>"",{c("result", 2)}="")'],
    fill=PatternFill("solid", fgColor=AMBER)))
for key in ("pl", "edge", "clv"):
    CF.add(f"{COL[key]}2:{COL[key]}{LAST}", CellIsRule(
        operator="greaterThan", formula=["0"],
        font=Font(name=FONT, size=10, bold=True, color="0B6B34")))
    CF.add(f"{COL[key]}2:{COL[key]}{LAST}", CellIsRule(
        operator="lessThan", formula=["0"],
        font=Font(name=FONT, size=10, bold=True, color="912018")))
for key in ("open_bal", "close_bal"):
    CF.add(f"{COL[key]}2:{COL[key]}{LAST}", FormulaRule(
        formula=[f'{c("date", 2)}=""'], font=Font(name=FONT, size=10, color=GREY)))
ws.auto_filter.ref = f"A1:{COL['notes']}{LAST}"

# ================================================================ DASHBOARD
wd.sheet_view.showGridLines = False
title_cell(wd, "B2", "BETTING LEDGER  -  DASHBOARD", size=18)
wd["B3"] = (f'=IF(COUNT({S}$B$3:$B$4)<2,'
            f'"SETTINGS INCOMPLETE - fill the two yellow cells on the Settings tab. '
            f'Until you do, every money figure below is wrong.","Settings complete.")')
wd["B3"].font = Font(name=FONT, size=10, bold=True, color="B54708")
wd.merge_cells("B3:M3")

# turnover counts what was actually at risk: the stake on a back, the liability on a lay
TURN = (f'SUMIFS({lg("size")},{lg("result")},"<>",{lg("side")},"<>Lay")'
        f'+SUMPRODUCT(({lg("result")}<>"")*({lg("side")}="Lay")'
        f'*{lg("size")}*({lg("got")}-1))')

blocks = [
    ("THE MONEY", "B", "C", [
        ("Bets logged", f'=COUNT({lg("date")})', INT),
        ("Bets settled", f'=COUNTA({lg("result")})', INT),
        ("Total staked", f'={TURN}', MONEY0),
        ("Net won / lost", f'=SUM({lg("pl")})', MONEY),
        ("Return on turnover", '=IFERROR($C$9/$C$8,"")', PCT),
        ("Balance now", f'={S}$B$3+$C$9', MONEY),
        ("Biggest drawdown", f'=MAX({lg("dd")})', MONEY0),
    ]),
    ("HOW IT IS GOING", "E", "F", [
        ("Strike rate",
         f'=IFERROR((COUNTIF({lg("result")},"Win")+COUNTIF({lg("result")},"Half Win"))'
         f'/$C$7,"")', PCT),
        ("Average odds taken", f'=IFERROR(AVERAGE({lg("got")}),"")', ODDS),
        ("Average bet size", f'=IFERROR(AVERAGE({lg("size")}),"")', MONEY),
        ("Biggest single win", f'=IFERROR(MAX({lg("pl")}),"")', MONEY),
        ("Biggest single loss", f'=IFERROR(MIN({lg("pl")}),"")', MONEY),
        ("Bets still open", '=$C$6-$C$7', INT),
    ]),
    ("IS YOUR PRICE JUDGEMENT GOOD?", "H", "I", [
        ("Average edge you claimed", f'=IFERROR(AVERAGE({lg("edge")}),"")', PCTS),
        ("Return you actually got", '=IFERROR($C$10,"")', PCTS),
        ("Gap between the two", '=IFERROR($I$6-$I$7,"")', PCTS),
        ("Bets taken UNDER your price", f'=COUNTIF({lg("edge")},"<0")', INT),
        ("What those cost you",
         f'=SUMIFS({lg("pl")},{lg("edge")},"<0")', MONEY),
        ("Verdict",
         f'=IF($C$7<20,"Too few settled bets to judge",'
         f'IF($I$8>0.1,"Odds Required looks too optimistic",'
         f'IF($I$8<-0.1,"Odds Required looks conservative","Broadly realistic")))', None),
    ]),
    ("CLOSING LINE VALUE", "K", "L", [
        ("Bets with a closing price", f'=COUNT({lg("close_od")})', INT),
        ("Average CLV %", f'=IFERROR(AVERAGE({lg("clv")}),"")', PCTS),
        ("Beat the close",
         f'=IFERROR(COUNTIF({lg("clv")},">0")/$L$6,"")', PCT),
        ("Still to fill in", '=$C$6-$L$6', INT),
    ]),
]
for title, lcol, vcol, rows in blocks:
    section(wd, f"{lcol}5", title)
    wd.merge_cells(f"{lcol}5:{vcol}5")
    for i, (lab, f, fmt) in enumerate(rows):
        r = 6 + i
        label(wd, f"{lcol}{r}", lab)
        value(wd, f"{vcol}{r}", f, fmt)
        wd[f"{vcol}{r}"].border = BOX
wd["I11"].alignment = Alignment(horizontal="center")
wd["I11"].font = Font(name=FONT, size=10, bold=True, color=NAVY)

label(wd, "B14",
      "Edge you claimed is the gap between Odds Required and the Betfair Odds you took. "
      "If that gap is much bigger than the return you actually got, your Odds Required is "
      "wishful and the price you 'need' is set too low.", italic=True, colour="667085")
wd.merge_cells("B14:M14")
wd["B14"].alignment = Alignment(wrap_text=True, vertical="top")
wd.row_dimensions[14].height = 30
label(wd, "B15",
      "CLV is the only measure here that does not depend on results. A bet can lose and "
      "still have beaten the close. Fill the Closing Odds column and this block comes alive.",
      italic=True, colour="667085")
wd.merge_cells("B15:M15")
wd["B15"].alignment = Alignment(wrap_text=True, vertical="top")
wd.row_dimensions[15].height = 26

for col, w in {"A": 3, "B": 27, "C": 15, "D": 3, "E": 24, "F": 14, "G": 3,
               "H": 28, "I": 16, "J": 3, "K": 25, "L": 14, "M": 3}.items():
    wd.column_dimensions[col].width = w
for ref in ("C9", "I8", "I10", "L7"):
    wd.conditional_formatting.add(ref, CellIsRule(
        operator="lessThan", formula=["0"],
        font=Font(name=FONT, size=11, bold=True, color="912018")))
for ref in ("C9", "L7"):
    wd.conditional_formatting.add(ref, CellIsRule(
        operator="greaterThan", formula=["0"],
        font=Font(name=FONT, size=11, bold=True, color="0B6B34")))

ch = LineChart()
ch.title = "Balance after every bet"
ch.style, ch.height, ch.width = 2, 10, 26
ch.y_axis.title, ch.x_axis.title = "Balance ($)", "Bet number"
ch.x_axis.delete = ch.y_axis.delete = False
ch.add_data(Reference(ws, min_col=IDX["close_bal"], min_row=1, max_row=LAST),
            titles_from_data=True)
ch.visible_cells_only = False
ch.display_blanks = "gap"
wd.add_chart(ch, "B17")

# ================================================================ BREAKDOWNS
wk.sheet_view.showGridLines = False
title_cell(wk, "A1", "BREAKDOWNS", size=18)
label(wk, "A2", "Where the money actually came from and went. A segment with only a "
                "handful of bets is noise, not a finding - read the Bets column first.",
      italic=True)

HC = ["Segment", "Bets", "Staked", "Won / Lost", "Return", "Strike rate",
      "Avg odds", "Avg CLV %"]
HF = [None, INT, MONEY0, MONEY, PCT, PCT, ODDS, PCTS]
tables = [("BY COMPETITION", "comp", "D", 26),
          ("BY BET TYPE", "bet_type", "E", 20),
          ("BY ODDS BAND", "band", "H", 7),
          ("BY MONTH", "month", "MONTH", 24)]

row, positions = 4, {}
for title, key, src, n in tables:
    cc = lg(key)
    section(wk, f"A{row}", title)
    wk.merge_cells(f"A{row}:H{row}")
    hr = row + 1
    for i, h in enumerate(HC):
        cell = wk.cell(row=hr, column=1 + i, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="475467")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BOX
    positions[key] = (hr, n)
    for j in range(n):
        r = hr + 1 + j
        if src == "MONTH":
            if j == 0:
                wk[f"K{r}"] = (f'=IF(COUNT({lg("date")})=0,"",'
                               f'DATE(YEAR(MIN({lg("date")})),MONTH(MIN({lg("date")})),1))')
            else:
                wk[f"K{r}"] = f'=IF($K{r-1}="","",EDATE($K{r-1},1))'
            wk[f"A{r}"] = f'=IF($K{r}="","",TEXT($K{r},"YYYY-MM"))'
        else:
            wk[f"A{r}"] = (f'=IF({S}${src}${L1+1+j}="","",{S}${src}${L1+1+j})')
        g = f'IF($A{r}="",""'
        wk[f"B{r}"] = f'={g},COUNTIFS({cc},$A{r}))'
        wk[f"C{r}"] = (f'={g},SUMIFS({lg("size")},{cc},$A{r},{lg("result")},"<>",'
                       f'{lg("side")},"<>Lay")+SUMPRODUCT(({cc}=$A{r})'
                       f'*({lg("result")}<>"")*({lg("side")}="Lay")*{lg("size")}'
                       f'*({lg("got")}-1)))')
        wk[f"D{r}"] = f'={g},SUMIFS({lg("pl")},{cc},$A{r}))'
        wk[f"E{r}"] = f'={g},IFERROR($D{r}/$C{r},""))'
        wk[f"F{r}"] = (f'={g},IFERROR((COUNTIFS({cc},$A{r},{lg("result")},"Win")'
                       f'+COUNTIFS({cc},$A{r},{lg("result")},"Half Win"))'
                       f'/COUNTIFS({cc},$A{r},{lg("result")},"<>"),""))')
        wk[f"G{r}"] = f'={g},IFERROR(AVERAGEIFS({lg("got")},{cc},$A{r}),""))'
        wk[f"H{r}"] = (f'={g},IFERROR(AVERAGEIFS({lg("clv")},{cc},$A{r},'
                       f'{lg("close_od")},">0"),""))')
        for i, fmt in enumerate(HF):
            cell = wk.cell(row=r, column=1 + i)
            cell.font = Font(name=FONT, size=10)
            cell.border = BOX
            if fmt:
                cell.number_format = fmt
    for col in ("D", "E", "H"):
        wk.conditional_formatting.add(f"{col}{hr+1}:{col}{hr+n}", CellIsRule(
            operator="lessThan", formula=["0"],
            font=Font(name=FONT, size=10, bold=True, color="912018")))
        wk.conditional_formatting.add(f"{col}{hr+1}:{col}{hr+n}", CellIsRule(
            operator="greaterThan", formula=["0"],
            font=Font(name=FONT, size=10, bold=True, color="0B6B34")))
    row = hr + n + 2

for col, w in {"A": 26, "B": 8, "C": 13, "D": 14, "E": 10,
               "F": 12, "G": 11, "H": 12, "K": 12}.items():
    wk.column_dimensions[col].width = w
wk.column_dimensions["K"].hidden = True
wk.freeze_panes = "B1"

hr, n = positions["bet_type"]
bar = BarChart()
bar.type, bar.title = "bar", "Won / lost by bet type"
bar.height, bar.width = 12, 20
bar.add_data(Reference(wk, min_col=4, min_row=hr, max_row=hr + n), titles_from_data=True)
bar.set_categories(Reference(wk, min_col=1, min_row=hr + 1, max_row=hr + n))
bar.visible_cells_only = False
wk.add_chart(bar, "J4")

# ================================================================ HOW TO USE
wh.sheet_view.showGridLines = False
title_cell(wh, "B2", "HOW TO USE THIS LEDGER", size=18)
steps = [
    ("SET IT UP - once", None),
    ("1.", "Save this file into OneDrive, not Documents on the C drive. That is what lets "
           "your phone open the same file."),
    ("2.", "Open the Settings tab along the bottom and fill the two yellow cells: your "
           "starting balance, and your Betfair commission rate (type 5% or 0.05)."),
    ("3.", "Still on Settings, add your own competitions and bet types on the blank rows "
           "under each list. The dropdowns pick them up straight away."),
    ("", ""),
    ("LOG A BET", None),
    ("4.", "Go to the Bet Log tab and the first empty row. Fill only the dark blue columns "
           "- the grey ones work themselves out."),
    ("5.", "Date: press Ctrl and the semicolon key together."),
    ("6.", "Odds Required is the price you decided you needed BEFORE you looked at the "
           "market. Betfair Odds is what you actually got. Fill both - the gap between "
           "them is the whole point."),
    ("7.", "Leave Back/Lay empty for a normal back bet. Only fill it in when you laid."),
    ("8.", "Win/Lose means did YOUR BET win, not did the team win. A lay that comes off "
           "is a Win. Half Win and Half Lose are for Asian lines."),
    ("9.", "Closing Odds is optional. Fill it when you have it and the CLV column and the "
           "Dashboard's CLV block come alive."),
    ("", ""),
    ("READ IT", None),
    ("10.", "Dashboard: the balance graph is the line you care about, and the third block "
            "answers the harder question - whether the price you say you need is one the "
            "market ever actually offers."),
    ("11.", "Breakdowns: where the money came from. Check the Bets column before believing "
            "anything - six bets is not a finding."),
    ("", ""),
    ("IF YOU NEED MORE ROOM", None),
    ("12.", "The log holds 1000 bets. To add more, click the row number of the last row, "
            "copy it, then paste into the rows beneath."),
]
r = 4
for num, text in steps:
    if text is None:
        section(wh, f"B{r}", num)
        wh.merge_cells(f"B{r}:F{r}")
        r += 1
        continue
    if num == "":
        r += 1
        continue
    label(wh, f"B{r}", num, bold=True, colour=TEAL)
    wh[f"B{r}"].alignment = Alignment(vertical="top")
    wh[f"C{r}"] = text
    wh[f"C{r}"].font = Font(name=FONT, size=10)
    wh[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    wh.merge_cells(f"C{r}:F{r}")
    wh.row_dimensions[r].height = max(15, 13 * (len(text) // 92 + 1))
    r += 1
r += 1
label(wh, f"B{r}",
      "One thing worth knowing: Total Staked counts your stake on a back bet and your "
      "liability on a lay, because that is what was really at risk. Return on turnover "
      "is measured against that, not against the stake box.", italic=True, colour="667085")
wh.merge_cells(f"B{r}:F{r}")
wh[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
wh.row_dimensions[r].height = 32
for col, w in {"A": 3, "B": 8, "C": 46, "D": 20, "E": 20, "F": 20}.items():
    wh.column_dimensions[col].width = w

for sheet in (ws, wd, wk, wset, wh):
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:1"

# ================================================================ demo
if "--demo" in sys.argv:
    import random
    from datetime import date, timedelta
    random.seed(5)
    wset["B3"], wset["B4"] = 2000, 0.05
    comps = ["EPL", "Serie A", "Liga MX", "Brasileirao", "Chinese Super League", "MLS"]
    types = ["Match Odds", "Over/Under 2.5 Goals", "Double Chance",
             "Both Teams To Score", "Asian Handicap"]
    d0 = date(2026, 8, 1)
    for i in range(60):
        r = 2 + i
        req = round(random.uniform(1.7, 4.5), 2)
        got = round(req * (1 + random.gauss(0.02, 0.06)), 2)
        p = 1 / got
        res = "Win" if random.random() < p else "Lose"
        if random.random() < 0.04:
            res = "Void"
        vals = {"date": d0 + timedelta(days=i), "comp": random.choice(comps),
                "teams": f"Team {i+1} v Team {i+40}", "bet_type": random.choice(types),
                "req": req, "got": got, "size": random.choice([25, 25, 50, 50, 75]),
                "result": res,
                "close_od": round(got * (1 + random.gauss(-0.01, 0.05)), 2)}
        if i >= 57:
            vals["result"], vals["close_od"] = "", ""
        for key, v in vals.items():
            cell = ws[f"{COL[key]}{r}"]
            cell.value = v if v != "" else None
            cell.font = Font(name=FONT, size=10)
            if key == "date":
                cell.number_format = DATEF

wb.active = 0
out = sys.argv[1] if len(sys.argv) > 1 else "Betting Ledger.xlsx"
wb.save(out)
print(f"written: {out}")
